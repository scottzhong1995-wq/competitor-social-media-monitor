from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def workbook(data):
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name, headers, rows):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            clean = [', '.join(map(str, x)) if isinstance(x, list) else ('不可用' if x is None else x) for x in row]
            ws.append(clean)
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    # Untrusted titles/contacts must never become spreadsheet formulas.
                    cell.data_type = 's'
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in ws[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='183D3A')
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for col in range(1, len(headers)+1):
            ws.column_dimensions[get_column_letter(col)].width = min(60, max(16, max(len(str(ws.cell(r, col).value or '')) for r in range(1, min(ws.max_row, 100)+1))+2))

    s = data['summary']
    sheet('口径与覆盖', ['项目', '说明'], [
        ['统计周', data['week']+' 至 '+data['end']], ['统计时区', data['timezone']],
        ['报告状态', '完整周' if data['complete'] else '进行中，非完整周'],
        ['数据截止', data['generated_at']], ['美国英语已确认新合作视频', s['confirmed_videos']],
        ['合作账号数', s['confirmed_creators']], ['已知累计播放', s['known_views']],
        ['播放量覆盖', f"{s['measured_videos']}/{s['confirmed_videos']} 条视频"],
        ['累计播放含义', '本周新发布、已确认合作、美国英语视频在周末前最后一次快照的播放量；非去重触达'],
        ['新增播放', '仅提供两次真实观测间的增量及起止时间；没有基线则留空，不推算整周'],
        ['历史覆盖', '目标近90天；首次采集无法恢复过去逐日播放量；导入/查询结果不代表全网总量'],
        ['独立内容', '仅在人工提供 content_group 时跨平台合并，未提供时按独立发布计'],
        ['搜索趋势', '美国 Google Web 相对热度0–100；同 series_id 同 query 相邻完整周比较；非绝对搜索量'],
        ['爆款候选', '按平台内内容格式的绝对播放阈值筛选；不代表因果分析或保证自然流量'],
        *[[k, v] for k, v in data['coverage'].items() if k != 'video_records']])
    fields = ['platform','title','url','creator','brands','models','format','published_at','discovered_at',
              'cooperation','evidence','market','language','direction','keywords','views','observed_at',
              'observed_growth','growth_from','growth_to','viral_candidate','viral_reason','source','metric_source','content_group']
    headers = ['平台','标题','视频链接','红人','品牌','车型','长短格式','发布时间','首次发现',
               '合作状态','合作证据','市场','语言','内容方向','标题关键词','累计播放','快照时间',
               '观测区间增量','区间起点','区间终点','高播放候选','筛选原因','视频来源','指标来源','跨平台内容组']
    sheet('本周发布视频', headers, [[v.get(k) for k in fields] for v in data['videos'] if v['new_in_week']])
    sheet('历史视频与效果', headers, [[v.get(k) for k in fields] for v in data['videos']])
    sheet('长短视频汇总', ['平台','格式','确认合作视频','已知累计播放','有播放数据条数'],
          [[r[k] for k in ['platform','format','videos','views','measured']] for r in data['breakdown']])
    fields = ['platform','creator','creator_url','followers','observed_at','creator_type','contact','contact_source','videos','confirmed']
    sheet('红人名单', ['平台','账号','主页','粉丝数','数据时间','红人类型','公开商务联系方式','联系方式来源','收录视频','确认合作视频'],
          [[r[k] for k in fields] for r in data['creators']])
    sheet('搜索趋势', ['同批次系列ID','搜索词','周起始','相对热度','周环比百分比','来源'],
          [[t[k] for k in ['series_id','query','week','value','change_percent','source']] for t in data['trends']])
    sheet('周度汇总', ['指标','数值'], [['数据标签',data.get('data_label','真实记录；覆盖受限')],['统计范围',data['scope']],['快照截止',data['snapshot_cutoff']],
          *[[k,v] for k,v in data['summary'].items()],['口径','多标签占比可超过100%；跨平台播放量未去重受众']])
    sheet('品牌平台车型', ['品牌','平台','车型','视频数','账号数','播放合计','均值','中位数'],
          [[r[k] for k in ('brand','platform','model','videos','creators','views','average_views','median_views')] for r in data['brand_breakdown']])
    sheet('爆款分析', ['原视频','标题','品牌','车型','红人','平台','发布时间','播放','高播放','相对爆款','基线样本','同龄中位数','倍数','置信度','增长每天','依据','分析材料','事实与推测','Qronge行动'],
          [[v['url'],v['title'],v['brands'],v['models'],v['creator'],v['platform'],v['published_at'],v['views'],v['high_views'],v['relative_viral'],
            v['baseline_sample_count'],v['baseline_median'],v['baseline_multiple'],v['baseline_confidence'],v['growth_per_day'],v['viral_reason'],v['analysis_basis'],
            str(v['materials']),v['action_suggestion']] for v in data['videos'] if v['viral_candidate']])
    sheet('效果与观察期', ['视频','形式','时长秒','分享','互动率百分比','公式','7天状态','7天播放','7天快照','30天状态','30天播放','30天快照','是否过期','可用状态'],
          [[v['url'],v['content_form'],v['duration_seconds'],v['shares'],v['interaction_rate'],v['interaction_formula'],v['age7']['status'],v['age7']['views'],v['age7']['observed_at'],
            v['age30']['status'],v['age30']['views'],v['age30']['observed_at'],v['stale'],v['availability']] for v in data['videos']])
    sheet('趋势口径与查询', ['系列','词','地区','渠道','范围','匹配','周','本周','上周','环比','部分周','采集时间','相关查询'],
          [[t[k] for k in ('series_id','query','geo','search_type','range_label','match_type','week','value','previous','change_label','partial','collected_at')]+[str(t['related']) if t['related'] else '不可用'] for t in data['trends']])
    sheet('红人扩展档案', ['平台','ID','名称','国家','所在地依据','语言','身份组','身份依据','首次观察合作','最近观察合作','基线样本','中位播放','爆款数','关注','备注','经纪信息','经纪来源'],
          [[c.get(k) for k in ('platform','account_id','creator','market','market_evidence','language','identity_group','identity_evidence','first_cooperation','last_cooperation','baseline_sample_count','baseline_median','viral_count','followed','notes','agency','agency_source')] for c in data['creators']])
    sheet('内容方向变化', ['方向','数量','占比百分比','对上周变化百分点'],[[r[k] for k in ('name','count','share','share_change_pp')] for r in data['directions']])
    sheet('模块分析与总结', ['模块','数据事实','分析与总结','Qronge行动（待验证）','限制','依据链接'],
          [[v['title'],'\n'.join(v['facts']),v['conclusion'],'\n'.join(v['actions']),'\n'.join(v['limitations']),'\n'.join(e['label']+' '+e['url'] for e in v['evidence'])] for v in data['insights'].values()])
    sheet('爆款深度分析', ['视频','原标题','候选类型','上线范围','合作状态','国家','语言','入选事实','内容观察','待核验推测','材料状态与来源','总结','Qronge行动','分析限制'],
          [[v['url'],v['title'],v['analysis']['kind'],v['analysis']['period'],v['cooperation'],v['market'],v['language'],
            '\n'.join(v['analysis']['facts']),'\n'.join(v['analysis']['observations']),'\n'.join(v['analysis']['hypotheses']) or '无附加推测',
            '\n'.join(m['label']+'：'+m['status']+' '+m['note']+' '+m['source'] for m in v['analysis']['materials']),
            v['analysis']['conclusion'],'\n'.join(v['analysis']['actions']),'\n'.join(v['analysis']['limitations'])] for v in data['viral_candidates']])
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def markdown(data):
    def value(v): return '不可用' if v is None else str(v)
    def bullets(items): return ['- '+item for item in items]
    s=data['summary']
    lines=[f"# Qronge 竞品红人合作周报 · {data['week']} 至 {data['end']}",
      f"数据标签：{data.get('data_label','真实收录记录，非全网统计')}。时区：{data['timezone']}；生成时间：{data['generated_at']}；快照截止：{data['snapshot_cutoff']}。",
      f"合作统计范围：{data['scope']}，过滤条件：{data['filters']}。完整周：{'是' if data['complete'] else '否'}。",
      '## 本周关键指标',
      f"- 发布 {s['published']} 条；本周首次发现 {s['discovered']} 条；有新快照 {s['updated']} 条。历史补录不计新上线。",
      f"- 确认合作 {s['confirmed_videos']} 条、{s['confirmed_creators']} 个平台账号；疑似 {s['pending']} 条和官方 {s['official']} 条单列，不相加为确认合作。",
      f"- 新上线内容已知累计播放 {value(s['known_views'])}；覆盖 {s['measured_videos']}/{s['published']}；历史内容观测窗口增量 {value(s['historical_growth'])}。跨平台播放未去重受众。",
      '## 各模块分析与总结',
      '以下为可追溯的规则分析与人工材料记录归纳，和看板、Excel共用同一结果；行动建议待验证，不代表已证实的传播原因。']
    for insight in data['insights'].values():
        lines += ['### '+insight['title'], '**数据事实**', *bullets(insight['facts']),
                  '**分析与总结**：'+insight['conclusion'], '**Qronge行动建议（待验证）**', *bullets(insight['actions']),
                  '**口径与限制**', *bullets(insight['limitations'])]
        if insight['evidence']:
            lines += ['**依据**']+[f"- [{e['label']}]({e['url']})" for e in insight['evidence']]
    lines += ['## 爆款视频清单与逐条分析',
      '本清单覆盖当前筛选视频列表中的非排除候选，含其他/未知市场及普通提及；合作状态逐条标明，不自动计入确认合作。高播放与相对账号爆款分别展示，历史视频不计本周新上线。']
    candidates=data['viral_candidates']
    lines.append(f"共 {len(candidates)} 条候选，正文展示前 {min(20,len(candidates))} 条；先按相对候选、基线倍数、播放量排序，完整清单及分析见Excel。")
    for v in candidates[:20]:
        a=v['analysis']
        lines += [f"### [{v['title']}]({v['url']})",
          f"**{a['period']} / {a['kind']}** · {'、'.join(v['brands'])} / {'、'.join(v['models']) or '车型未知'} · {v['platform']} / {v['content_form']} · {v['creator']}",
          f"发布：{v['published_at']}；合作状态：{v['cooperation']}；国家/语言：{v['market']}/{v['language']}；[合作证据]({v['evidence_source']})：{v['evidence'] or '尚未确认'}。",
          '**入选依据与表现**',*bullets(a['facts']),
          '**内容观察**',*bullets(a['observations']),a['basis'],
          '**封面、开场、测试场景、结构、观众讨论与字幕**']
        for m in a['materials']:
            lines.append(f"- {m['label']}：{m['status']}"+(f"；[{m['kind']}] {m['note']}；[材料来源]({m['source']})" if m['source'] else '。'))
        lines += ['**推测（待核验）**',*bullets(a['hypotheses'] or ['无附加推测；未取得的材料不补写。']),
                  '**分析与总结**：'+a['conclusion'], '**可供Qronge验证的选题与合作行动**',*bullets(a['actions']),
                  '**判定限制**',*bullets(a['limitations'])]
    if not candidates: lines.append('暂无达到当前规则的已收录候选；不代表竞品不存在爆款。')
    lines+=['## 视频列表合作证据与来源（含不同合作状态）']+[f"- [{v['title']}]({v['url']})：{v['cooperation']}；{v['evidence'] or '未发现证据'}；[证据]({v['evidence_source']})" for v in data['videos'] if v['new_in_week']]
    lines+=['## 数据限制']+bullets(data['limitations'])
    return '\n\n'.join(lines)+'\n'
