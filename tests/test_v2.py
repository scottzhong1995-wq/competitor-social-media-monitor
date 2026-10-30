import io
import json
import os
import tempfile
import unittest
from datetime import timedelta
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from monitor.core import connect,import_videos,report,review,week_bounds,import_trends
from monitor.profiles import upsert_account
from monitor.settings import load,save
from monitor.export import workbook,markdown


class V2Tests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.db=connect(Path(self.temp.name)/'test.db')
        self.start,self.end=week_bounds()

    def tearDown(self):
        self.db.close();self.temp.cleanup()

    def item(self, **kw):
        v=dict(platform='YouTube',url='https://youtu.be/target00001',title='Yozma IN-10 review',creator='test',creator_id='stable-id',
               published_at=(self.start+timedelta(hours=1)).isoformat(),observed_at=(self.start+timedelta(days=1)).isoformat(),
               source='isolated test',format='long',market='US',language='en',cooperation='明确合作',evidence='paid disclosure',views=1000)
        v.update(kw);return v

    def test_official_and_suspected_do_not_enter_default_cooperation(self):
        import_videos(self.db,[self.item(cooperation='品牌官方内容'),self.item(url='https://youtu.be/target00002',cooperation='疑似合作')])
        r=report(self.db);self.assertEqual(r['summary']['confirmed_videos'],0)
        self.assertEqual(r['summary']['official'],1);self.assertEqual(r['summary']['pending'],1)
        self.assertEqual(report(self.db,scope='official')['summary']['published'],1)

    def test_baseline_same_age_and_no_ratio_when_insufficient(self):
        target_pub=self.start-timedelta(days=10)
        for i in range(10):
            pub=target_pub-timedelta(days=(i+1)*2)
            import_videos(self.db,[self.item(url=f'https://youtu.be/base{i:07d}',published_at=pub.isoformat(),observed_at=(pub+timedelta(days=7)).isoformat(),views=100,baseline_only=True)])
        import_videos(self.db,[self.item(published_at=target_pub.isoformat(),observed_at=(target_pub+timedelta(days=7)).isoformat(),views=400)])
        v=report(self.db)['videos'][0]
        self.assertTrue(v['relative_viral']);self.assertFalse(v['high_views']);self.assertEqual(v['baseline_multiple'],4)
        with self.db:self.db.execute("UPDATE videos SET format='short' WHERE id='yt:base0000000'")
        v=report(self.db)['videos'][0]
        self.assertIsNone(v['baseline_multiple']);self.assertEqual(v['baseline_sample_count'],9)

    def test_no_same_age_snapshot_means_no_precise_multiple(self):
        import_videos(self.db,[self.item(published_at=(self.start-timedelta(days=15)).isoformat(),views=999999)])
        v=report(self.db)['videos'][0]
        self.assertTrue(v['high_views']);self.assertIsNone(v['baseline_multiple']);self.assertEqual(v['age7']['status'],'缺少同龄快照')

    def test_no_new_snapshot_means_no_weekly_growth(self):
        pub=self.start-timedelta(days=30)
        for d,n in [(1,20),(3,40)]:
            import_videos(self.db,[self.item(published_at=pub.isoformat(),observed_at=(pub+timedelta(days=d)).isoformat(),views=n)])
        self.assertIsNone(report(self.db)['videos'][0]['observed_growth'])

    def test_manual_audit_and_material_source_required(self):
        import_videos(self.db,[self.item()])
        with self.assertRaises(ValueError):review(self.db,{'id':'yt:target00001','materials':{'opening':{'note':'unseen'}}})
        review(self.db,{'id':'yt:target00001','cooperation':'普通提及','market':'unknown','review_note':'not sponsored'})
        self.assertEqual(self.db.execute('SELECT count(*) FROM reviews').fetchone()[0],1)
        import_videos(self.db,[self.item()])
        self.assertEqual(report(self.db)['videos'][0]['market'],'unknown')

    def test_creator_identity_requires_evidence_and_accounts_not_deleted(self):
        with self.assertRaises(ValueError):
            upsert_account(self.db,{'platform':'TikTok','account_id':'t','creator':'x','identity_group':'one'})
        with self.db:upsert_account(self.db,{'platform':'TikTok','account_id':'t','creator':'x','followed':True,'notes':'keep'})
        self.assertEqual(report(self.db)['creators'][0]['notes'],'keep')

    def test_zero_trend_and_series_scope(self):
        monday=self.start.date();earlier=monday-timedelta(days=7)
        header='series_id,query,week,value,source,geo,search_type\n'
        import_trends(self.db,header+f'a,Yozma,{earlier},0,test,US,web\na,Yozma,{monday},50,test,US,web')
        t=report(self.db)['trends'][-1]
        self.assertIsNone(t['change_percent']);self.assertEqual(t['change_label'],'从 0 增至 50')
        with self.assertRaises(ValueError):import_trends(self.db,header+f'a,Yozma,{monday},50,test,GB,web')

    def test_xlsx_and_chinese_report_share_filtered_summary(self):
        import_videos(self.db,[self.item(),self.item(url='https://youtu.be/target00002',title='Tuttio Soleil01 review',views=None)])
        r=report(self.db,brand='Yozma')
        wb=load_workbook(io.BytesIO(workbook(r)))
        exported={row[0]:row[1] for row in wb['周度汇总'].iter_rows(min_row=2,values_only=True)}
        self.assertEqual(exported['published'],r['summary']['published'])
        self.assertEqual(exported['known_views'],1000)
        self.assertIn('发布 1 条',markdown(r));self.assertEqual(len(r['videos']),1)

    def test_configurable_timezone_keywords_and_thresholds(self):
        with patch.dict(os.environ,{'MONITOR_CONFIG':str(Path(self.temp.name)/'settings.json')}):
            c=load();c['report_timezone']='Asia/Shanghai';c['brands'].append({'name':'Newbrand','models':[],'website':'https://example.com'});save(c)
            from monitor.core import identify
            self.assertEqual(identify('Newbrand review')[0],['Newbrand'])
            self.assertEqual(week_bounds('2026-10-26')[0].hour,16)
