import io
import json
import os
import socket
import subprocess
import sys
import tempfile
import time
import unittest
from datetime import timedelta
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from openpyxl import load_workbook
from monitor.core import ROOT, week_bounds


class WebTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        with socket.socket() as s:
            s.bind(('127.0.0.1', 0))
            port = s.getsockname()[1]
        cls.base = f'http://127.0.0.1:{port}'
        cls.process = subprocess.Popen([sys.executable, '-m', 'monitor', 'serve', '--port', str(port)],
            cwd=ROOT, env={**os.environ, 'MONITOR_DB':cls.temp.name+'/web.db'}, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        try:
            for _ in range(100):
                try:
                    with urlopen(cls.base+'/health', timeout=1) as r:
                        if r.status == 200:
                            break
                except OSError:
                    if cls.process.poll() is not None:
                        raise RuntimeError('测试服务器启动失败')
                    time.sleep(.05)
            else:
                raise RuntimeError('测试服务器未就绪')
            with urlopen(cls.base+'/api/config') as r:
                cls.token = json.load(r)['csrf']
        except Exception:
            cls.process.terminate(); cls.process.wait(); cls.temp.cleanup()
            raise

    @classmethod
    def tearDownClass(cls):
        cls.process.terminate(); cls.process.wait(timeout=5); cls.temp.cleanup()

    def test_page_empty_report_and_export(self):
        with urlopen(self.base) as r:
            self.assertIn('竞品雷达', r.read().decode())
        with urlopen(self.base+'/api/report') as r:
            data = json.load(r)
            self.assertIn('summary', data)
        with urlopen(self.base+'/api/export') as r:
            wb = load_workbook(io.BytesIO(r.read()))
            self.assertIn('红人名单', wb.sheetnames)

    def test_mutations_require_csrf_and_invalid_dates_fail(self):
        with self.assertRaises(HTTPError) as error:
            urlopen(Request(self.base+'/api/import', data=b'[]', method='POST'))
        self.assertEqual(error.exception.code, 403)
        with self.assertRaises(HTTPError) as error:
            urlopen(self.base+'/api/report?week=2026-10-27')
        self.assertEqual(error.exception.code, 400)

    def test_import_review_and_report_functional_flow(self):
        start, _ = week_bounds()
        record = dict(platform='YouTube', url='https://youtu.be/testvideo01', title='Tuttio Soleil01 GB Pro review',
            creator='ISOLATED TEST', published_at=(start+timedelta(hours=1)).isoformat(),
            observed_at=(start+timedelta(days=1)).isoformat(), source='isolated web test', views=100000,
            market='US', language='en', format='short')
        def post(path, body):
            with urlopen(Request(self.base+path, data=json.dumps(body).encode(), headers={'X-Monitor-Token':self.token}, method='POST')) as r:
                return json.load(r)
        self.assertEqual(post('/api/import',[record])['imported'],1)
        post('/api/review',dict(id='yt:testvideo01', cooperation='明确合作', evidence='fixture disclosure', format='short',market='US',language='en'))
        with urlopen(self.base+'/api/report') as r:
            report = json.load(r)
        self.assertEqual(report['summary']['confirmed_videos'],1)
        self.assertEqual(report['breakdown'][0]['format'],'short')
        self.assertTrue(report['videos'][0]['viral_candidate'])


if __name__ == '__main__':
    unittest.main()
