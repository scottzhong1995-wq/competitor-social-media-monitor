import io
import tempfile
import unittest
from datetime import timedelta
from pathlib import Path

from openpyxl import load_workbook
from monitor.core import connect,import_videos,import_trends,report,week_bounds
from monitor.export import workbook,markdown


class InsightTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.db=connect(Path(self.temp.name)/'test.db')
        self.start,self.end=week_bounds()

    def tearDown(self):
        self.db.close();self.temp.cleanup()

    def video(self,**kw):
        v=dict(platform='YouTube',url='https://youtu.be/insight0001',title='Yozma IN-10 Pro range test',
            creator='fixture',creator_id='fixture-id',published_at=(self.start+timedelta(hours=1)).isoformat(),
            observed_at=(self.end-timedelta(hours=1)).isoformat(),source='isolated test',format='long',market='US',language='en',
            cooperation='明确合作',evidence='test disclosure',views=80000)
        v.update(kw);return v

    def test_empty_modules_do_not_invent_campaign_or_viral(self):
        r=report(self.db)
        self.assertEqual(len(r['insights']),9)
        self.assertEqual(r['viral_candidates'],[])
        self.assertIn('不能解读为竞品没有',r['insights']['overview']['conclusion'])
        self.assertIn('缺少',r['insights']['overview']['conclusion'])
        self.assertIn('暂无达标',markdown(r))

    def test_high_views_does_not_invent_relative_multiple_or_materials(self):
        import_videos(self.db,[self.video()])
        v=report(self.db)['viral_candidates'][0];a=v['analysis']
        self.assertEqual(a['kind'],'高播放视频')
        self.assertIn('只满足绝对高播放',a['conclusion'])
        self.assertEqual(a['hypotheses'],[])
        self.assertTrue(all(not m['source'] for m in a['materials']))
        self.assertIn('续航对比',' '.join(a['actions']))
        self.assertTrue(any('增长速度不可用' in f for f in a['facts']))

    def test_observations_and_hypotheses_are_separate_and_sourced(self):
        import_videos(self.db,[self.video(materials={
            'opening':{'note':'TEST前5秒列出测试条件','source':'https://example.invalid/opening','kind':'观察'},
            'comments':{'note':'TEST受众可能关注价格','source':'https://example.invalid/comments','kind':'推测'}})])
        r=report(self.db);a=r['viral_candidates'][0]['analysis']
        self.assertIn('TEST前5秒',' '.join(a['observations']))
        self.assertNotIn('TEST受众可能',' '.join(a['observations']))
        self.assertIn('TEST受众可能',' '.join(a['hypotheses']))
        doc=markdown(r)
        self.assertNotIn('https://example.invalid/opening',doc)
        self.assertNotIn('推测（待核验）',doc)
        self.assertNotIn('封面：不可用',doc)

    def test_shared_summary_and_viral_analysis_in_excel_and_report(self):
        import_videos(self.db,[self.video()])
        r=report(self.db);wb=load_workbook(io.BytesIO(workbook(r)))
        signals=[row[0] for row in wb['关键信号'].iter_rows(min_row=2,values_only=True)]
        self.assertEqual(signals,r['weekly_signals'])
        self.assertTrue(all(x in markdown(r) for x in signals))
        headers=[c.value for c in wb['爆款视频'][1]]
        self.assertEqual(headers,['视频','标题','数据','关键词','内容方向'])
        self.assertNotIn('爆款深度分析',wb.sheetnames)
        self.assertNotIn('各模块分析与总结',markdown(r))

    def test_current_filters_apply_to_analysis_and_search_trends(self):
        import_videos(self.db,[self.video(),self.video(url='https://youtu.be/insight0002',title='Tuttio Soleil01 review')])
        header='series_id,query,week,value,source,geo,search_type\n'
        day=self.start.date()
        import_trends(self.db,header+f'a,Yozma,{day},40,test,US,web\na,Tuttio,{day},70,test,US,web')
        r=report(self.db,brand='Yozma')
        self.assertNotIn('Tuttio',' '.join(r['insights']['brands']['facts']))
        self.assertEqual([t['query'] for t in r['trends']],['Yozma'])
        self.assertEqual(len(r['viral_candidates']),1)

    def test_historical_candidate_and_ordinary_mention_are_labelled(self):
        import_videos(self.db,[self.video(published_at=(self.start-timedelta(days=10)).isoformat(),cooperation='普通提及')])
        r=report(self.db,scope='all');v=r['viral_candidates'][0]
        self.assertEqual(r['summary']['published'],0)
        self.assertEqual(v['analysis']['period'],'历史内容复查')
        self.assertIn('普通提及',' '.join(v['analysis']['facts']))
        self.assertEqual(r['summary']['confirmed_videos'],0)

    def test_relative_candidate_has_same_age_evidence(self):
        pub=self.start-timedelta(days=10)
        for i in range(10):
            older=pub-timedelta(days=2*(i+1))
            import_videos(self.db,[self.video(url=f'https://youtu.be/older{i:06d}',published_at=older.isoformat(),observed_at=(older+timedelta(days=7)).isoformat(),views=1000,baseline_only=True)])
        import_videos(self.db,[self.video(published_at=pub.isoformat(),observed_at=(pub+timedelta(days=7)).isoformat(),views=4000)])
        a=report(self.db)['viral_candidates'][0]['analysis']
        self.assertEqual(a['kind'],'相对爆款候选')
        self.assertIn('基线 4 倍',' '.join(a['facts']))
        self.assertIn('10 条可比样本',' '.join(a['facts']))

    def test_previous_week_comparison_uses_local_week_across_dst(self):
        # Los Angeles Nov 2 Monday is UTC08; the preceding Monday is UTC07.
        # A fixed 168h subtraction would wrongly exclude this first half-hour.
        import_videos(self.db,[self.video(published_at='2026-10-26T07:30:00Z',observed_at='2026-10-26T08:00:00Z')])
        r=report(self.db,'2026-11-02')
        self.assertEqual(r['previous_week_start'],'2026-10-26T07:00:00+00:00')
        self.assertEqual(r['summary']['previous_videos'],1)
        self.assertIn('Yozma -1 条',' '.join(r['insights']['overview']['facts']))
