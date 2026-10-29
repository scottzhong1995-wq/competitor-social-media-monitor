import io
import json
import os
import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from monitor.core import (canonical, connect, identify, import_trends, import_videos,
                          report, review, stamp, week_bounds)
from monitor.export import workbook
from monitor.__main__ import due_jobs
from monitor.youtube import collect, api


class MonitorTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.db = connect(Path(self.temp.name)/'test.db')
        self.start, self.end = week_bounds()
        self.week = self.start.astimezone(ZoneInfo('America/Los_Angeles')).date().isoformat()

    def tearDown(self):
        self.db.close()
        self.temp.cleanup()

    def video(self, **kw):
        v = dict(platform='YouTube', url='https://www.youtube.com/watch?v=abcdefghijk',
                 title='Yozma IN-10 Pro honest review', creator='Test creator',
                 description='Sponsored by Yozma. Test fixture only.',
                 published_at=(self.start+timedelta(hours=1)).isoformat(),
                 observed_at=(self.start+timedelta(days=1)).isoformat(),
                 source='test-fixture', views=100, format='long', market='US', language='en')
        v.update(kw)
        return v

    def test_specific_models_do_not_double_count_base(self):
        self.assertEqual(identify('Yozma IN-10 Pro')[1], ['IN-10 Pro'])
        self.assertEqual(identify('Tuttio Soleil01 GB Pro')[1], ['Soleil01 GB Pro'])
        self.assertEqual(identify('IN-10 Pro vs IN-10')[1], ['IN-10 Pro','IN-10'])

    def test_dedupe_and_snapshots_do_not_rewrite_discovery(self):
        import_videos(self.db, [self.video()])
        discovered = self.db.execute('SELECT discovered_at FROM videos').fetchone()[0]
        import_videos(self.db, [self.video(url='https://youtu.be/abcdefghijk?utm_source=test',
                      observed_at=(self.start+timedelta(days=2)).isoformat(), views=150)])
        self.assertEqual(self.db.execute('SELECT count(*) FROM videos').fetchone()[0], 1)
        self.assertEqual(self.db.execute('SELECT count(*) FROM snapshots').fetchone()[0], 2)
        self.assertEqual(self.db.execute('SELECT discovered_at FROM videos').fetchone()[0], discovered)
        row = report(self.db, self.week)['videos'][0]
        self.assertEqual(row['observed_growth'], 50)

    def test_missing_baseline_is_not_zero_growth(self):
        import_videos(self.db, [self.video()])
        self.assertIsNone(report(self.db)['videos'][0]['observed_growth'])

    def test_confirmed_cooperation_requires_evidence_and_us(self):
        import_videos(self.db, [self.video(market='unknown')])
        self.assertEqual(report(self.db)['summary']['confirmed_videos'], 0)
        with self.assertRaises(ValueError):
            review(self.db, dict(id='yt:abcdefghijk', cooperation='明确合作', evidence='', format='long'))
        review(self.db, dict(id='yt:abcdefghijk', cooperation='明确合作', evidence='paid sponsorship disclosure in description', format='long', market='US', language='en'))
        self.assertEqual(report(self.db)['summary']['confirmed_videos'], 1)
        import_videos(self.db, [self.video(market='unknown', format='unknown')])
        self.assertEqual(report(self.db)['videos'][0]['cooperation'], '明确合作')
        self.assertEqual(report(self.db)['videos'][0]['format'], 'long')

    def test_no_future_metric_leak_into_past_week(self):
        import_videos(self.db, [self.video()])
        import_videos(self.db, [self.video(observed_at=(self.end+timedelta(hours=1)).isoformat(), views=900000)])
        historical = report(self.db, self.week)['videos'][0]
        self.assertEqual(historical['views'], 100)
        self.assertFalse(historical['viral_candidate'])

    def test_unknown_views_not_zero(self):
        import_videos(self.db, [self.video(views=None, cooperation='明确合作', evidence='disclosure')])
        r = report(self.db)
        self.assertIsNone(r['summary']['known_views'])
        self.assertEqual(r['summary']['measured_videos'], 0)
        self.assertIsNone(r['breakdown'][0]['views'])

    def test_negative_growth_is_flagged(self):
        import_videos(self.db, [self.video()])
        import_videos(self.db, [self.video(observed_at=(self.start+timedelta(days=2)).isoformat(), views=80)])
        self.assertTrue(report(self.db)['videos'][0]['metric_correction'])

    def test_atomic_validation_and_contact_source(self):
        with self.assertRaises(ValueError):
            import_videos(self.db, [self.video(), self.video(url='https://youtu.be/zzzzzzzzzzz', contact='test@example.com')])
        self.assertEqual(self.db.execute('SELECT count(*) FROM videos').fetchone()[0], 0)
        for value in (-1, 1.2, True):
            with self.assertRaises(ValueError):
                import_videos(self.db, [self.video(views=value)])

    def test_unrelated_and_invalid_urls_rejected(self):
        with self.assertRaises(ValueError):
            import_videos(self.db, [self.video(title='Nothing related', description='')])
        for url in ['https://youtube.com.evil.test/watch?v=abcdefghijk', 'javascript:alert(1)', 'http://youtu.be/abcdefghijk']:
            with self.assertRaises(ValueError):
                canonical('YouTube', url)
        self.assertEqual(canonical('TikTok', 'https://www.tiktok.com/@test/video/123456789')[0], 'tt:123456789')

    def test_ninety_day_window(self):
        import_videos(self.db, [self.video(published_at=(self.start-timedelta(days=100)).isoformat())])
        self.assertEqual(report(self.db)['videos'], [])

    def test_la_week_boundaries_include_dst(self):
        start, end = week_bounds('2026-10-26')
        self.assertEqual((end-start).total_seconds()/3600, 169)
        self.assertEqual(start.hour, 7)
        self.assertEqual(end.hour, 8)
        with self.assertRaises(ValueError):
            week_bounds('2026-10-27')

    def test_scheduler_beijing_night_and_tuesday_reporting(self):
        zone = 'Asia/Shanghai'
        for hour in range(24):
            now = datetime(2026,10,27,hour,tzinfo=ZoneInfo(zone))
            collection, export = due_jobs(now, {}, zone)
            self.assertEqual(collection, hour == 2)
            self.assertEqual(export, hour == 3)
        now = datetime(2026,10,27,2,tzinfo=ZoneInfo(zone))
        self.assertFalse(due_jobs(now, {'collection':'2026-10-27'}, zone)[0])
        monday = datetime(2026,10,26,3,tzinfo=ZoneInfo(zone))
        self.assertFalse(due_jobs(monday, {}, zone)[1])
        # At Tuesday 03:00 Beijing, the preceding US week really has ended.
        self.assertTrue(week_bounds(now=now)[1] < now)

    def test_trends_series_isolation_and_zero_baseline(self):
        first = (self.start.astimezone(ZoneInfo('America/Los_Angeles')).date()-timedelta(days=7)).isoformat()
        header = 'series_id,query,week,value,source,geo,search_type\n'
        text = header+f'a,Yozma,{first},20,test,US,web\na,Yozma,{self.week},40,test,US,web\nb,Yozma,{self.week},99,test,US,web\n'
        self.assertEqual(import_trends(self.db, text), 3)
        rows = report(self.db)['trends']
        self.assertEqual(next(r for r in rows if r['series_id']=='a' and r['week']==self.week)['change_percent'],100)
        self.assertIsNone(next(r for r in rows if r['series_id']=='b')['change_percent'])
        with self.assertRaises(ValueError):
            import_trends(self.db, header+f'c,Yozma,{first},101,test,US,web')

    def test_excel_openable_with_literal_untrusted_titles(self):
        import_videos(self.db, [self.video(title='=HYPERLINK("https://evil.invalid","Yozma IN-10")')])
        wb = load_workbook(io.BytesIO(workbook(report(self.db))), data_only=False)
        self.assertEqual(len(wb.sheetnames), 6)
        cell = wb['本周发布视频']['B2']
        self.assertEqual(cell.data_type, 's')
        self.assertTrue(cell.value.startswith('='))

    def test_youtube_adapter_through_mocked_api(self):
        video = self.video()
        def fake(resource, **kwargs):
            if resource == 'search':
                return {'items': [{'id': {'videoId':'abcdefghijk'}}]}
            if resource == 'channels':
                return {'items':[{'id':'channel', 'statistics':{'subscriberCount':'123'}}]}
            return {'items':[dict(id='abcdefghijk', snippet=dict(title=video['title'], description='Sponsored test fixture', channelId='channel', channelTitle='Test account', publishedAt=video['published_at']), statistics={'viewCount':'555'})]}
        with patch('monitor.youtube.api', side_effect=fake):
            collect(self.db, backfill=True)
        self.assertEqual(self.db.execute('SELECT count(*) FROM videos').fetchone()[0],1)
        self.assertEqual(self.db.execute('SELECT views FROM snapshots').fetchone()[0],555)
        self.assertEqual(self.db.execute('SELECT market FROM videos').fetchone()[0],'unknown')
        self.assertEqual(self.db.execute('SELECT format FROM videos').fetchone()[0],'unknown')

    def test_missing_api_key_reported_without_secrets(self):
        with patch.dict(os.environ, {}, clear=True):
            with self.assertRaises(RuntimeError):
                collect(self.db)
        run = self.db.execute('SELECT status,detail FROM runs').fetchone()
        self.assertEqual(run['status'], 'failed')
        self.assertIn('YOUTUBE_API_KEY', run['detail'])

    def test_api_key_sent_in_header_not_url_and_error_redacted(self):
        with patch.dict(os.environ, {'YOUTUBE_API_KEY':'test-only-secret'}):
            with patch('monitor.youtube.urlopen', side_effect=OSError('test-only-secret')) as opener:
                with self.assertRaises(RuntimeError) as error:
                    api('search', q='Yozma')
        request = opener.call_args.args[0]
        self.assertNotIn('test-only-secret', request.full_url)
        self.assertEqual(request.get_header('X-goog-api-key'),'test-only-secret')
        self.assertNotIn('test-only-secret', str(error.exception))


if __name__ == '__main__':
    unittest.main()
